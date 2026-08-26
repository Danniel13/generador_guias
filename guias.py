from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx2pdf import convert

EMU_PER_INCH = 914400


def set_table_no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement("w:tblBorders")

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        borders.append(elem)

    tblPr.append(borders)


def fit_size_keep_aspect(img_path, max_w_in, max_h_in):
    with Image.open(img_path) as im:
        w_px, h_px = im.size

    ar = w_px / h_px

    w1 = max_w_in
    h1 = w1 / ar

    if h1 <= max_h_in:
        return w1, h1

    h2 = max_h_in
    w2 = h2 * ar

    return w2, h2


def configurar_documento(doc, margin_in):
    """
    Configura el documento en orientación horizontal
    y establece los márgenes.
    """

    sec = doc.sections[0]

    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)

    sec.left_margin = Inches(margin_in)
    sec.right_margin = Inches(margin_in)
    sec.top_margin = Inches(margin_in)
    sec.bottom_margin = Inches(margin_in)


def agregar_imagen_al_documento(doc, img, margin_in, gap_in):
    """
    Agrega una nueva página al documento con la misma
    estructura que tenía el archivo individual original.
    """

    sec = doc.sections[0]

    page_w = sec.page_width / EMU_PER_INCH
    page_h = sec.page_height / EMU_PER_INCH

    avail_w = page_w - (2 * margin_in)
    avail_h = page_h - (2 * margin_in)

    cell_w = (avail_w - gap_in) / 2
    cell_h = avail_h

    # Crear tabla de 1 fila y 2 columnas
    table = doc.add_table(rows=1, cols=2)

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    set_table_no_borders(table)

    # Colocar la misma imagen dos veces,
    # igual que en la lógica original
    for c in range(2):
        cell = table.cell(0, c)

        cell.width = Inches(cell_w)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        w_fit, h_fit = fit_size_keep_aspect(
            str(img),
            cell_w,
            cell_h
        )

        run = p.add_run()

        run.add_picture(
            str(img),
            width=Inches(w_fit),
            height=Inches(h_fit)
        )


def main():
    print("Seleccione el tipo de salida:")
    print("1 -> Word (.docx)")
    print("2 -> PDF")

    while True:
        try:
            tipo = int(input("Ingrese opción (1 o 2): "))

            if tipo in [1, 2]:
                break

            print("Opción inválida.")

        except ValueError:
            print("Ingrese un número válido.")

    images_dir = Path("./")
    output_dir = Path("./salida")

    output_dir.mkdir(parents=True, exist_ok=True)

    valid_ext = {
        ".jpg",
        ".jpeg",
        ".png",
        ".bmp",
        ".webp",
        ".tif",
        ".tiff"
    }

    imgs = sorted(images_dir.glob("*.*"))

    # Filtrar únicamente imágenes válidas
    imagenes_validas = [
        img for img in imgs
        if img.suffix.lower() in valid_ext
    ]

    if not imagenes_validas:
        print("No se encontraron imágenes válidas.")
        return

    # ==================================================
    # CREAR UN ÚNICO DOCUMENTO
    # ==================================================

    doc = Document()

    configurar_documento(
        doc=doc,
        margin_in=0.3
    )

    total = len(imagenes_validas)

    for numero, img in enumerate(imagenes_validas, start=1):

        print(f"Procesando {numero}/{total}: {img.name}")

        # Agregar el contenido de esta imagen
        agregar_imagen_al_documento(
            doc=doc,
            img=img,
            margin_in=0.3,
            gap_in=0.2
        )

        # Agregar salto de página si todavía
        # faltan imágenes por procesar
        if numero < total:
            doc.add_page_break()

    # ==================================================
    # GUARDAR UN ÚNICO ARCHIVO
    # ==================================================

    output_docx = output_dir / "guias.docx"

    doc.save(output_docx)

    print("\nProceso completado.")
    print(f"DOCX creado: {output_docx}")

    # ==================================================
    # SI EL USUARIO ELIGIÓ PDF
    # ==================================================

    if tipo == 2:

        output_pdf = output_dir / "guias.pdf"

        convert(
            str(output_docx),
            str(output_pdf)
        )

        print(f"PDF creado: {output_pdf}")

        # Eliminar el DOCX temporal
        output_docx.unlink(missing_ok=True)

        print(f"DOCX temporal eliminado: {output_docx}")


if __name__ == "__main__":
    main()